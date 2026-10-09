# -*- coding: utf-8 -*-
# Usage: python3 build_ep.py <output.pdf> <data_module>
# Builds an intermediate .docx in a temp dir, converts it to PDF with LibreOffice, keeps only the PDF.
import re, sys, os, shutil, subprocess, tempfile, importlib
from collections import Counter, defaultdict
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
d = importlib.import_module(sys.argv[2])
M = d.META
TURNS = 150

KO = "WenQuanYi Zen Hei"; LAT = "Carlito"
NAVY = RGBColor(0x1B,0x3A,0x5C); GRAY = RGBColor(0x66,0x66,0x66)
SPK = defaultdict(lambda: RGBColor(0x55,0x55,0x55), {
    "Inès":RGBColor(0xA2,0x2C,0x4F), "Mathieu":RGBColor(0x1F,0x5A,0x8C),
    "Hélène":RGBColor(0x8A,0x4B,0x0F), "Baptiste":RGBColor(0x4B,0x6B,0x2E), "Sofia":RGBColor(0x6A,0x3D,0x9A)})
LVL = {"A2":"E3F1E3","B1":"DCEBF7","B2":"FFF0D2","C1":"F6DAE4"}

def fr(t):
    return re.sub(r" ([?!:;»])", " \\1", t.replace("« ", "« "))

def setfont(run, size=None, bold=None, italic=None, color=None):
    run.font.name = LAT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    rf.set(qn("w:eastAsia"), KO); rf.set(qn("w:ascii"), LAT); rf.set(qn("w:hAnsi"), LAT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = color

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.0)
st = doc.styles["Normal"]; st.font.name = LAT; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), KO)

def para(text="", size=10.5, bold=False, italic=False, color=None, align=None, after=4, before=0, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = 1.2
    if keep: pf.keep_with_next = True
    if align is not None: p.alignment = align
    if text: setfont(p.add_run(text), size, bold, italic, color)
    return p

def heading(text):
    p = para(after=6, before=14, keep=True)
    setfont(p.add_run(text), 16, True, None, NAVY)
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k,v in (("val","single"),("sz","6"),("space","1"),("color","1B3A5C")): bt.set(qn("w:"+k), v)
    b.append(bt); pPr.append(b)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),hexcolor); tcPr.append(sh)

def fixw(t, widths):
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
        for c in t.columns[i].cells: c.width = Cm(w)

def fill(cell, text, w, bold=False, color=None, size=9.5):
    cell.width = Cm(w); cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    setfont(p.add_run(text), size, bold, None, color)

def table(headers, rows, widths, level_col=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i,h in enumerate(headers):
        fill(t.rows[0].cells[i], h, widths[i], True, RGBColor(255,255,255)); shade(t.rows[0].cells[i], "1B3A5C")
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"),"true"); t.rows[0]._tr.get_or_add_trPr().append(th)
    for r in rows:
        cells = t.add_row().cells
        cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"),"true"); t.rows[-1]._tr.get_or_add_trPr().append(cant)
        for i,txt in enumerate(r):
            if i == 1: txt = fr(txt)
            fill(cells[i], txt, widths[i], bold=(i==1))
            if level_col is not None and i == level_col: shade(cells[i], LVL[txt])
    fixw(t, widths)

def line(s, text, size):
    p = para(after=3)
    p.paragraph_format.left_indent = Cm(2.4); p.paragraph_format.first_line_indent = Cm(-2.4)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(2.4))
    setfont(p.add_run(s + " :"), 10.5, True, None, SPK[s]); setfont(p.add_run("\t" + text), size)

# footer page number
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
setfont(fp.add_run(f"Conseil de Paris · Partie {M['part']} · Épisode {M['ep']} · "), 8.5, None, None, GRAY)
r = fp.add_run(); setfont(r, 8.5, None, None, GRAY)
for kind, txt in (("begin",None),("instr"," PAGE "),("end",None)):
    if kind == "instr":
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"),"preserve"); e.text = txt
    else:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
    r._r.append(e)

# ---------- Title block ----------
para(M["part_label"], 11, True, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER, 2, 20)
para(f"제{M['ep']}화  {M['title_fr']}", 24, True, False, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 2)
para(M["title_ko"], 13, False, True, GRAY, WD_ALIGN_PARAGRAPH.CENTER, 14)

cnt = Counter(x[0] for x in d.D if x[0])
turns = sum(cnt.values())
assert turns == TURNS, f"{turns} turns, expected {TURNS}"
cs = " · ".join(f"{k} {v}" for k, v in sorted(cnt.items(), key=lambda kv: -kv[1]))
meta = [("장소·시간", M["place"]), ("등장인물", M["chars"]), ("업무 주제", M["work"]), ("일상·로맨스", M["life"]),
        ("학습 포인트", M["focus"]), ("분량·난이도", f"대화 {TURNS}턴 ({cs}) · {M['levelmix']}")]
t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"; t.autofit = False
for k,v in meta:
    c = t.add_row().cells
    fill(c[0], k, 3.0, True); fill(c[1], v, 14.0); shade(c[0], "E3EAF2")
fixw(t, [3.0, 14.0])

# ---------- Dialogue ----------
heading("1. 대화문 Dialogue")
para("※ 턴 번호 없이 화자 이름으로 구분합니다. 이탤릭 지문은 턴 수에 포함되지 않습니다. 해석은 문서 맨 끝(5장)에 있습니다.", 9, False, True, GRAY, after=8)
for s, f, k in d.D:
    if s is None: para(fr(f), 9.5, False, True, GRAY, after=6, before=4)
    else: line(s, fr(f), 11)

# ---------- Vocab ----------
heading("2. 주요 단어·표현 분석")
para("레벨 색상: B1 파랑 · B2 노랑 · C1 분홍", 9, False, True, GRAY, after=6)
table(["레벨","표현","뜻","용법·메모"], d.VOCAB, [1.3,4.2,4.3,7.2], level_col=0)

# ---------- Grammar ----------
heading("3. 문법 분석")
for lv, title, body in d.GRAM:
    p = para(after=2, before=8, keep=True)
    setfont(p.add_run(f"[{lv}] "), 10.5, True, None, NAVY); setfont(p.add_run(title), 11, True)
    p2 = para(after=4); p2.paragraph_format.left_indent = Cm(0.6); setfont(p2.add_run(fr(body)), 10)

# ---------- Culture ----------
heading("4. 문화·업계 메모")
for title, body in d.CULTURE:
    p = para(after=1, before=6, keep=True); setfont(p.add_run("• " + title), 10.5, True)
    p2 = para(after=3); p2.paragraph_format.left_indent = Cm(0.6); setfont(p2.add_run(body), 10)

# ---------- Translation (last) ----------
doc.add_page_break()
heading("5. 한글 해석")
para("대화 순서대로 수록했습니다.", 9, False, True, GRAY, after=8)
for s, f, k in d.D:
    if s is None: para(k, 9.5, False, True, GRAY, after=6, before=4)
    else: line(s, k, 10.5)

para("다음 화 예고 — " + M["next"], 10, True, False, NAVY, after=0, before=18)

# ---------- Export PDF only ----------
out = os.path.abspath(sys.argv[1])
with tempfile.TemporaryDirectory() as tmp:
    base = os.path.splitext(os.path.basename(out))[0]
    src = os.path.join(tmp, base + ".docx"); doc.save(src)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    shutil.move(os.path.join(tmp, base + ".pdf"), out)
print(out, turns, "turns")
