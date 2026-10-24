# -*- coding: utf-8 -*-
"""Combine one part's episode PDFs into a single book PDF with cover, contents and bookmarks.
usage: make_book.py <series_dir> <part_no> <out.pdf> <workdir>"""
import sys, glob, os, subprocess, importlib.util
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

series_dir, part, out, work = sys.argv[1], int(sys.argv[2]), sys.argv[3], sys.argv[4]
eps = []
for f in sorted(glob.glob(f"{series_dir}/tools/part{part}/ep*_data.py")):
    sp = importlib.util.spec_from_file_location("m", f); m = importlib.util.module_from_spec(sp); sp.loader.exec_module(m)
    n = m.META["ep"]
    pdf = glob.glob(f"{series_dir}/part{part}/ep{n:02d}_*.pdf")
    assert len(pdf) == 1, (n, pdf)
    eps.append((n, m.META, pdf[0], len(PdfReader(pdf[0]).pages)))
M0 = eps[0][1]
NAVY = RGBColor(0x1F, 0x3A, 0x5F); GRAY = RGBColor(0x66, 0x66, 0x66)

def build_front(start_pages):
    doc = Document(); s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.2)
    def p(text, size, bold=False, italic=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6):
        para = doc.add_paragraph(); para.alignment = align
        para.paragraph_format.space_before = Pt(before); para.paragraph_format.space_after = Pt(after)
        r = para.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.name = "Calibri"; r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
        if color: r.font.color.rgb = color
        return para
    # cover
    p(M0["series"], 30, True, color=NAVY, before=150)
    p(M0["part_label"], 15, color=GRAY, after=40)
    p(f"제{part}부 · 제{eps[0][0]}화 – 제{eps[-1][0]}화", 13, after=4)
    p(f"{len(eps)}화 · 화당 120턴 · 총 {120*len(eps):,}턴", 11, color=GRAY, after=4)
    LANGS = {"fr": "프랑스어", "es": "스페인어", "it": "이탈리아어", "de": "독일어", "nl": "네덜란드어", "ja": "일본어"}
    p(f"비즈니스 {LANGS.get(M0.get('lang', 'fr'), '')} 연속 대화문 · 한글 해석 포함", 11, italic=True, color=GRAY)
    doc.add_page_break()
    # contents
    p("목차 · Índice", 18, True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, after=14)
    t = doc.add_table(rows=0, cols=4); t.autofit = False
    for (n, M, _, _), sp_ in zip(eps, start_pages):
        c = t.add_row().cells
        for cell, txt, w, b, col in ((c[0], f"제{n}화", 1.8, True, NAVY), (c[1], M["title_fr"], 7.4, True, None),
                                      (c[2], M["title_ko"], 5.2, False, GRAY), (c[3], f"p. {sp_}", 1.6, False, None)):
            cell.width = Cm(w); para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(5)
            if cell is c[3]: para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = para.add_run(txt); r.font.size = Pt(11); r.bold = b; r.font.name = "Calibri"
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
            if col: r.font.color.rgb = col
    docx = os.path.join(work, "front.docx"); doc.save(docx)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{work}/lo", "--headless", "--convert-to", "pdf", "--outdir", work, docx],
                   check=True, capture_output=True)
    return os.path.join(work, "front.pdf")

# two passes: front-matter length decides the episode start pages
front_len = 2
for _ in range(2):
    starts, cur = [], front_len + 1
    for e in eps: starts.append(cur); cur += e[3]
    fp = build_front(starts)
    n = len(PdfReader(fp).pages)
    if n == front_len: break
    front_len = n

w = PdfWriter()
for pg in PdfReader(fp).pages: w.add_page(pg)
for (n, M, pdf, cnt), sp_ in zip(eps, starts):
    for pg in PdfReader(pdf).pages: w.add_page(pg)
    w.add_outline_item(f"제{n}화 {M['title_fr']} · {M['title_ko']}", sp_ - 1)
w.add_metadata({"/Title": f"{M0['series']} – {M0['part_label']}"})
with open(out, "wb") as fh: w.write(fh)
print(out, len(w.pages), "pages")
