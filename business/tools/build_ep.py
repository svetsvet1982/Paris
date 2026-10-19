# -*- coding: utf-8 -*-
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import importlib, importlib.util, os
# argv[2]: module name (e.g. part1.ep01_data) or path to a data .py file
if sys.argv[2].endswith(".py"):
    _spec = importlib.util.spec_from_file_location("epdata", sys.argv[2]); d = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(d)
else:
    d = importlib.import_module(sys.argv[2])
M = d.META
LANG = M.get("lang", "fr")

KO = "Malgun Gothic"; LAT = "Calibri"
WINE = RGBColor(0x1F,0x3A,0x5F); GRAY = RGBColor(0x66,0x66,0x66)
from collections import defaultdict
# speaker colors in order of first appearance
_PAL = [RGBColor(0x2E,0x6B,0x3A),RGBColor(0x1F,0x5A,0x8C),RGBColor(0x8A,0x4B,0x0F),RGBColor(0x7A,0x1F,0x3D),RGBColor(0x0F,0x6E,0x6E),RGBColor(0x55,0x3C,0x7A),RGBColor(0x6B,0x5B,0x1A),RGBColor(0x3D,0x3D,0x3D)]
SPK = defaultdict(lambda: RGBColor(0x55,0x3C,0x7A))
for _s in dict.fromkeys(x[0] for x in d.D if x[0] and x[0] != "#"):
    SPK[_s] = _PAL[len(SPK) % len(_PAL)]
# CEFR levels, or JLPT levels for Japanese
LVL = {"A1":"E3F1E3","A2":"E3F1E3","B1":"DCEBF7","B2":"FFF0D2","C1":"F6DAE4","C2":"F6DAE4",
       "N5":"E3F1E3","N4":"E3F1E3","N3":"DCEBF7","N2":"FFF0D2","N1":"F6DAE4"}
LEGEND = M.get("legend", "레벨 색상: A2 초록 · B1 파랑 · B2 노랑 · C1 분홍")
COLON = " :" if LANG == "fr" else ":"
SERIES = M.get("series", "Maison Delcourt")
# east-asian font for target-language text (Japanese dialogue needs a JP font)
DLG_EA = "IPAGothic" if LANG == "ja" else "Malgun Gothic"

def fr(t):
    if LANG != "fr": return t
    t = t.replace("« ", "« ")
    t = re.sub(r" ([?!:;»])", " \\1", t)
    return t

def setfont(run, size=None, bold=None, italic=None, color=None, ko=KO, lat=LAT):
    run.font.name = lat
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    rf.set(qn("w:eastAsia"), ko); rf.set(qn("w:ascii"), lat); rf.set(qn("w:hAnsi"), lat)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = color

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)

for name in ("Normal",):
    st = doc.styles[name]; st.font.name = LAT; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), KO)
for name, size in (("Heading 1",16),("Heading 2",13)):
    st = doc.styles[name]; st.font.name = LAT; st.font.size = Pt(size); st.font.bold = True
    st.font.color.rgb = WINE
    st.element.rPr.rFonts.set(qn("w:eastAsia"), KO)
    st.element.rPr.rFonts.set(qn("w:ascii"), LAT); st.element.rPr.rFonts.set(qn("w:hAnsi"), LAT)

def para(text="", size=10.5, bold=False, italic=False, color=None, align=None, after=4, before=0, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = 1.2
    if keep: pf.keep_with_next = True
    if align: p.alignment = align
    if text:
        setfont(p.add_run(text), size, bold, italic, color)
    return p

def heading(text, level=1):
    p = doc.add_heading(level=level)
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    setfont(p.add_run(text), 16 if level==1 else 13, True, None, WINE)
    if level == 1:
        pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
        for k,v in (("val","single"),("sz","6"),("space","1"),("color","1F3A5F")): bt.set(qn("w:"+k), v)
        b.append(bt); pPr.append(b)
    return p

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),hexcolor); tcPr.append(sh)

def fixw(t, widths):
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
        for c in t.columns[i].cells: c.width = Cm(w)

def table(headers, rows, widths, level_col=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    def fill(cell, text, w, bold=False, color=None, size=9.5):
        cell.width = Cm(w); cell.text = ""
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        setfont(p.add_run(text), size, bold, None, color)
    for i,h in enumerate(headers):
        fill(t.rows[0].cells[i], h, widths[i], True, RGBColor(255,255,255)); shade(t.rows[0].cells[i], "1F3A5F")
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"),"true"); trPr.append(th)
    for r in rows:
        cells = t.add_row().cells
        cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"),"true"); t.rows[-1]._tr.get_or_add_trPr().append(cant)
        for i,txt in enumerate(r):
            if i == 1 and len(headers)>=4: txt = fr(txt)
            fill(cells[i], txt, widths[i], bold=(i==1 and len(headers)>=4))
            if level_col is not None and i == level_col: shade(cells[i], LVL[txt])
    fixw(t, widths)
    return t

# footer page number
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_ff = M.get("footer_fmt") or (SERIES + " · " + M.get("footer_part", "Partie") + " {part} · " + M.get("footer_ep", "Épisode") + " {ep} · ")
setfont(fp.add_run(_ff.format(series=SERIES, part=M["part"], ep=M["ep"])), 8.5, None, None, GRAY, ko=DLG_EA)
r = fp.add_run(); setfont(r, 8.5, None, None, GRAY)
for t_,txt in (("begin",None),("instr"," PAGE "),("end",None)):
    if t_=="instr":
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"),"preserve"); e.text = txt
    else:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t_)
    r._r.append(e)

# ---------- Title block ----------
para(M["part_label"], 11, True, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER, 2, 20)
para(f"제{M['ep']}화  {M['title_fr']}", 24, True, False, WINE, WD_ALIGN_PARAGRAPH.CENTER, 2)
para(M["title_ko"], 13, False, True, GRAY, WD_ALIGN_PARAGRAPH.CENTER, 14)

from collections import Counter
cnt = Counter(x[0] for x in d.D if x[0] and x[0] != "#")
turns = sum(cnt.values())
assert turns == M.get("turns", 120), turns
cs = " · ".join(f"{k} {v}" for k, v in sorted(cnt.items(), key=lambda kv: -kv[1]))
meta = [("장소·시간", M["place"]), ("등장인물", M["chars"]), ("비즈니스 주제", M["topic"]), ("학습 포인트", M["focus"])]
meta = M.get("info", meta) + [("분량·난이도", f"대화 {turns}턴 ({cs}) · {M['levelmix']}")]
t = doc.add_table(rows=0, cols=2); t.style="Table Grid"; t.autofit=False
for k,v in meta:
    c = t.add_row().cells
    for cell,txt,w,b in ((c[0],k,3.0,True),(c[1],v,14.0,False)):
        cell.width = Cm(w); cell.text=""; p = cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
        setfont(p.add_run(txt), 9.5, b)
    shade(c[0], "E6ECF3")
fixw(t, [3.0, 14.0])

# ---------- Sections ----------
# Optional d.PRE (before dialogue) and d.SECTIONS (after dialogue) let each series use its own layout:
#   ("table", title, note, headers, rows, widths, level_col)  ("notes", title, note, [(tag, title, body)])
#   ("bullets", title, note, [(title, body)])                  ("text", title, note, [paragraph, ...])
if hasattr(d, "SECTIONS"):
    POST = d.SECTIONS
else:
    POST = [("table", "핵심 비즈니스 표현", "이번 화에서 반복 연습할 기능별 표현입니다.", ["기능","표현","뜻","대화 속 예문"], d.KEY, [2.9,4.6,3.6,5.9], None),
            ("table", "주요 단어·표현 분석", LEGEND, ["레벨","표현","뜻","용법·메모"], d.VOCAB, [1.3,4.2,4.3,7.2], 0),
            ("notes", "문법 분석", None, d.GRAM),
            ("bullets", "비즈니스 문화 메모", None, d.CULTURE)]
PRE = getattr(d, "PRE", [])
NUM = iter(range(1, 100))

def section(sec):
    kind, title, note = sec[0], sec[1], sec[2]
    heading(f"{next(NUM)}. {title}")
    if note: para(note, 9, False, True, GRAY, after=6)
    if kind == "table":
        _, _, _, headers, rows, widths, lc = sec
        table(headers, [tuple(r) for r in rows], widths, level_col=lc)
    elif kind == "notes":
        for tag, ttl, body in sec[3]:
            p = para(after=2, before=8, keep=True)
            if tag: setfont(p.add_run(f"[{tag}] "), 10.5, True, None, WINE)
            setfont(p.add_run(ttl), 11, True)
            p2 = para(after=4); p2.paragraph_format.left_indent = Cm(0.6)
            setfont(p2.add_run(fr(body)), 10)
    elif kind == "bullets":
        for ttl, body in sec[3]:
            p = para(after=1, before=6, keep=True); setfont(p.add_run("• " + ttl), 10.5, True)
            p2 = para(after=3); p2.paragraph_format.left_indent = Cm(0.6); setfont(p2.add_run(fr(body)), 10)
    elif kind == "text":
        for ptxt in sec[3]:
            setfont(para(after=4).add_run(fr(ptxt)), 10)

for sec in PRE: section(sec)

# ---------- Dialogue ----------
# speaker "#" = scene heading (not a turn), None = stage direction (not a turn)
heading(f"{next(NUM)}. " + M.get("dialogue_title", "대화문 Dialogue"))
para(f"※ 턴 번호 없이 화자 이름으로 구분합니다. 이탤릭 지문은 턴 수에 포함되지 않습니다. 해석은 문서 맨 끝({len(PRE) + len(POST) + 2}장)에 있습니다.", 9, False, True, GRAY, after=8)
for s, f, k in d.D:
    if s == "#":
        setfont(para(after=4, before=10, keep=True).add_run(fr(f)), 10.5, True, None, WINE, ko=DLG_EA)
    elif s is None:
        setfont(para(after=6, before=4).add_run(fr(f)), 9.5, False, True, GRAY, ko=DLG_EA)
    else:
        p = para(after=3)
        p.paragraph_format.left_indent = Cm(2.4); p.paragraph_format.first_line_indent = Cm(-2.4)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(2.4))
        setfont(p.add_run(s + COLON), 10.5, True, None, SPK[s]); setfont(p.add_run("\t" + fr(f)), 11, ko=DLG_EA)

for sec in POST: section(sec)

# ---------- Translation (last) ----------
doc.add_page_break()
heading(f"{next(NUM)}. 한글 해석")
para("대화 순서대로 수록했습니다.", 9, False, True, GRAY, after=8)
for s, f, k in d.D:
    if s == "#":
        para(k, 10.5, True, False, WINE, after=4, before=10, keep=True)
    elif s is None:
        para(k, 9.5, False, True, GRAY, after=6, before=4)
    else:
        p = para(after=3)
        p.paragraph_format.left_indent = Cm(2.4); p.paragraph_format.first_line_indent = Cm(-2.4)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(2.4))
        setfont(p.add_run(s + COLON), 10.5, True, None, SPK[s]); setfont(p.add_run("\t" + k), 10.5)

para(M.get("next_label", "다음 화 예고") + " — " + M["next"], 10, True, False, WINE, after=0, before=18)

doc.save(sys.argv[1])
