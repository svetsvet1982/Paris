# -*- coding: utf-8 -*-
"""Combine all parts of a series into one PDF: cover, contents by part, part dividers, nested bookmarks.
usage: make_full_book.py <series_dir> <out.pdf> <workdir>"""
import sys, glob, os, subprocess, importlib.util
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

series_dir, out, work = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(work, exist_ok=True)
parts = []  # (part_no, part_label, [(ep, META, pdf, pages)])
for pn in range(1, 10):
    files = sorted(glob.glob(f"{series_dir}/tools/part{pn}/ep*_data.py"))
    if not files: break
    eps = []
    for f in files:
        sp = importlib.util.spec_from_file_location("m", f); m = importlib.util.module_from_spec(sp); sp.loader.exec_module(m)
        n = m.META["ep"]; pdf = glob.glob(f"{series_dir}/part{pn}/ep{n:02d}_*.pdf"); assert len(pdf) == 1, (n, pdf)
        eps.append((n, m.META, pdf[0], len(PdfReader(pdf[0]).pages)))
    parts.append((pn, eps[0][1]["part_label"], eps))
M0 = parts[0][2][0][1]
NAVY = RGBColor(0x1F, 0x3A, 0x5F); GRAY = RGBColor(0x66, 0x66, 0x66)
LANGS = {"fr": "프랑스어", "es": "스페인어", "it": "이탈리아어", "de": "독일어", "nl": "네덜란드어", "ja": "일본어"}
total_eps = sum(len(p[2]) for p in parts)

def mkdoc():
    d = Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7); s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.2)
    return d
def para(doc, text, size, bold=False, italic=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6):
    p = doc.add_paragraph(); p.alignment = align; p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.name = "Calibri"
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
    if color: r.font.color.rgb = color
    return p
def to_pdf(doc, name):
    path = os.path.join(work, name + ".docx"); doc.save(path)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{work}/lo", "--headless", "--convert-to", "pdf", "--outdir", work, path], check=True, capture_output=True)
    return os.path.join(work, name + ".pdf")

def front(starts):
    doc = mkdoc()
    para(doc, M0["series"], 32, True, color=NAVY, before=150)
    para(doc, f"비즈니스 {LANGS.get(M0.get('lang'), '')} 연속 대화문 · 전권", 14, color=GRAY, after=40)
    para(doc, f"{len(parts)}부 · {total_eps}화 · 화당 120턴 · 총 {120 * total_eps:,}턴", 12, after=4)
    para(doc, "한글 해석 포함", 11, italic=True, color=GRAY)
    doc.add_page_break()
    para(doc, "목차", 18, True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, after=10)
    for (pn, label, eps), (pstart, epstarts) in zip(parts, starts):
        para(doc, f"{label}   p. {pstart}", 11.5, True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=2)
        for (n, M, _, _), st in zip(eps, epstarts):
            para(doc, f"   제{n}화  {M['title_fr']} · {M['title_ko']}  ·····  p. {st}", 9.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    return to_pdf(doc, "front")

def divider(pn, label, eps):
    doc = mkdoc()
    para(doc, f"제{pn}부", 26, True, color=NAVY, before=220)
    para(doc, label, 15, color=GRAY, after=10)
    para(doc, f"제{eps[0][0]}화 – 제{eps[-1][0]}화", 11, color=GRAY)
    return to_pdf(doc, f"div{pn}")

divs = [divider(pn, label, eps) for pn, label, eps in parts]
front_len = 3
for _ in range(3):
    cur = front_len + 1; starts = []
    for (pn, label, eps), dv in zip(parts, divs):
        pstart = cur; cur += len(PdfReader(dv).pages); epst = []
        for e in eps: epst.append(cur); cur += e[3]
        starts.append((pstart, epst))
    fp = front(starts); n = len(PdfReader(fp).pages)
    if n == front_len: break
    front_len = n

w = PdfWriter()
for pg in PdfReader(fp).pages: w.add_page(pg)
for (pn, label, eps), dv, (pstart, epst) in zip(parts, divs, starts):
    for pg in PdfReader(dv).pages: w.add_page(pg)
    parent = w.add_outline_item(label, pstart - 1)
    for (n, M, pdf, cnt), st in zip(eps, epst):
        for pg in PdfReader(pdf).pages: w.add_page(pg)
        w.add_outline_item(f"제{n}화 {M['title_fr']} · {M['title_ko']}", st - 1, parent=parent)
w.add_metadata({"/Title": f"{M0['series']} – 전권"})
w.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
with open(out, "wb") as fh: w.write(fh)
print(out, len(w.pages), "pages", round(os.path.getsize(out) / 1e6, 1), "MB")
